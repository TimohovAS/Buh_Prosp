"""Выезд: оплата работника за день, командировочные и сверка начислений с выплатами."""

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import func, select

from backend.models import Client, Income, User, WorkDiaryWorkerDay, Worker
from backend.routers.work_diaries_router import (
    create_entry,
    create_invoice_from_entries,
    delete_entry,
    export_proposal_xlsx,
    get_project_costs,
    get_summary,
    get_worker_days,
    list_entries,
    update_entry,
    update_payout_reconciliation,
)
from backend.routers.workers_router import _calculate_payout, create_worker_payout, update_worker_payout_link
from backend.work_diary_costing import allocate
from backend.schemas import (
    WorkDiaryEntryCreate,
    WorkDiaryEntryUpdate,
    WorkDiaryInvoiceCreate,
    WorkDiaryInvoiceLineCreate,
    WorkDiaryMaterialCreate,
    WorkDiaryPayoutReconciliationUpdate,
    WorkDiaryProposalExportRequest,
    WorkDiaryWorkerDayInput,
    WorkerPayoutCreate,
    WorkerPayoutLinkUpdate,
)

DAY = date(2026, 9, 1)
NEXT_DAY = date(2026, 9, 2)


def _user(db, name):
    user = User(username=name, password_hash="hash", role="admin")
    db.add(user)
    return user


async def _worker(db, name="Ibrahim", **rates):
    values = {"regular_day_rate": Decimal("6000"), "billing_hourly_rate": Decimal("1500")}
    values.update(
        {key: Decimal(str(value)) if isinstance(value, int | float) else value for key, value in rates.items()}
    )
    worker = Worker(name=name, **values)
    db.add(worker)
    await db.flush()
    return worker


async def _entry(db, user, project, workers, *, day=DAY, hours=6, pay_mode=None, **fields):
    worker_days = fields.pop("worker_days", None)
    if worker_days is None and pay_mode is not None:
        worker_days = [WorkDiaryWorkerDayInput(worker_id=worker.id, pay_mode=pay_mode) for worker in workers]
    return await create_entry(
        WorkDiaryEntryCreate(
            date=day,
            project_id=project.id,
            worker_ids=[worker.id for worker in workers],
            description=fields.pop("description", "Radovi na objektu"),
            duration_hours=hours,
            worker_days=worker_days,
            **fields,
        ),
        db,
        user,
    )


async def _costs(db, user, project, date_from=None, date_to=None):
    return await get_project_costs(project.id, date_from, date_to, db, user)


async def _worker_day_count(db) -> int:
    return (await db.execute(select(func.count(WorkDiaryWorkerDay.id)))).scalar_one()


# --- 1–2. Полный день: себестоимость — дневная ставка, заказчику — часы на объекте ---


async def test_full_day_costs_the_day_rate_but_bills_only_hours_on_site(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-full-day")
    worker = await _worker(db_session)

    hourly = await _entry(db_session, user, project, [worker], day=date(2026, 8, 31))
    full_day = await _entry(db_session, user, project, [worker], pay_mode="full_day", is_trip=True, travel_hours=3)

    # По часам 6 ч стоят 6 × 750; тот же выезд с оплатой полного дня — вся дневная ставка.
    assert hourly.labor_amount == 4500
    assert full_day.duration_hours == 6
    assert full_day.person_hours == 6
    assert full_day.labor_amount == 6000
    assert full_day.day_labor_amount == 6000
    assert full_day.hourly_labor_amount == 0
    assert full_day.total_cost_amount == 6000
    # Заказчику — только фактические 6 часов, оплата полного дня сумму не повышает.
    assert full_day.billable_amount == hourly.billable_amount == 6 * 1500
    assert full_day.margin_amount == 9000 - 6000
    assert full_day.is_trip is True
    assert full_day.travel_hours == 3
    pay = full_day.worker_pay[0]
    assert (pay.pay_mode, pay.day_rate, pay.share, pay.rate_missing, pay.day_rate_manual) == (
        "full_day",
        6000,
        1,
        False,
        False,
    )


async def test_trip_mark_alone_does_not_switch_pay_to_a_full_day(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-mark-only")
    worker = await _worker(db_session)

    entry = await _entry(db_session, user, project, [worker], is_trip=True, travel_hours=3, travel_km=200)

    assert entry.is_trip is True
    assert entry.worker_pay[0].pay_mode == "hourly"
    assert entry.labor_amount == 4500
    assert await _worker_day_count(db_session) == 0


async def test_exit_service_raises_the_customer_amount_but_not_the_cost(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-exit-service")
    worker = await _worker(db_session)

    entry = await _entry(
        db_session,
        user,
        project,
        [worker],
        pay_mode="full_day",
        materials=[
            WorkDiaryMaterialCreate(
                description="Troškovi izlaska",
                source="service",
                quantity=1,
                unit="usl",
                unit_price_snapshot=2500,
            )
        ],
    )

    assert entry.billable_service_amount == 2500
    assert entry.billable_amount == 6 * 1500 + 2500
    assert entry.labor_amount == 6000
    assert entry.total_cost_amount == 6000
    assert entry.margin_amount == 11500 - 6000


# --- 3. Две записи одного работника за день: ставка и суточные не удваиваются ---


async def test_day_rate_and_per_diem_are_charged_once_per_worker_day(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-two-entries")
    worker = await _worker(
        db_session,
        "Denis",
        trip_pricing_mode="allowances",
        trip_work_day_rate=8000,
        trip_per_diem_rate=2500,
        trip_food_rate=3000,
    )

    first = await _entry(db_session, user, project, [worker], hours=4, pay_mode="trip_day")
    assert first.labor_amount == 8000
    assert first.allowance_amount == 2500 + 3000

    # Вторая запись того же дня снова отмечает суточные — строка дня одна на работника.
    second = await _entry(
        db_session,
        user,
        project,
        [worker],
        hours=2,
        worker_days=[WorkDiaryWorkerDayInput(worker_id=worker.id, pay_mode="trip_day", per_diem_amount=2500)],
    )
    entries = {entry.id: entry for entry in await list_entries(project.id, None, None, None, db_session, user)}
    first, second = entries[first.id], entries[second.id]

    # Доли в парах: 8000 × 4/6 вниз до пары, остаток — последней записи дня.
    assert (first.labor_amount, second.labor_amount) == (5333.33, 2666.67)
    assert (first.allowance_amount, second.allowance_amount) == (3666.66, 1833.34)
    assert first.worker_pay[0].day_entries_count == 2
    assert first.worker_pay[0].day_hours == 6
    assert await _worker_day_count(db_session) == 1

    summary = await get_summary(project.id, None, None, None, db_session, user)
    assert summary.labor_amount == 8000
    assert summary.allowance_amount == 5500
    costs = await _costs(db_session, user, project)
    assert costs.labor_amount == 8000
    assert costs.allowance_amount == 5500
    assert costs.total_cost_amount == 13500


# --- 4. Два проекта за день: дневная ставка делится по часам ---


async def test_full_day_is_split_between_projects_by_hours(db_session, make_project):
    museum = await make_project(db_session, code="PR-A", name="Muzej")
    fryer = await make_project(db_session, code="PR-B", name="Friteza")
    user = _user(db_session, "trip-two-projects")
    worker = await _worker(db_session)

    first = await _entry(db_session, user, museum, [worker], hours=1.5, pay_mode="full_day")
    # Режим оплаты — на работника и дату: вторая запись этого дня делит тот же день.
    second = await _entry(db_session, user, fryer, [worker], hours=4.5)

    assert (await _costs(db_session, user, museum)).labor_amount == pytest.approx(1500)
    assert (await _costs(db_session, user, fryer)).labor_amount == pytest.approx(4500)
    assert second.worker_pay[0].pay_mode == "full_day"
    assert second.worker_pay[0].share == pytest.approx(0.75)

    await update_entry(second.id, WorkDiaryEntryUpdate(duration_hours=1.5), db_session, user)
    assert (await _costs(db_session, user, museum)).labor_amount == pytest.approx(3000)
    assert (await _costs(db_session, user, fryer)).labor_amount == pytest.approx(3000)

    # Перенос записи на другую дату освобождает день: вся ставка остаётся первой записи.
    moved = await update_entry(second.id, WorkDiaryEntryUpdate(date=NEXT_DAY), db_session, user)
    assert moved.worker_pay[0].pay_mode == "hourly"
    assert (await _costs(db_session, user, museum)).labor_amount == pytest.approx(6000)

    await update_entry(moved.id, WorkDiaryEntryUpdate(date=DAY), db_session, user)
    await delete_entry(moved.id, db_session, user)
    assert (await _costs(db_session, user, museum)).labor_amount == pytest.approx(6000)
    assert await _worker_day_count(db_session) == 1

    await delete_entry(first.id, db_session, user)
    assert await _worker_day_count(db_session) == 0


def test_allocation_is_in_paras_and_never_negative():
    assert allocate(Decimal("100"), [Decimal("2")] * 3) == [Decimal("33.33"), Decimal("33.33"), Decimal("33.34")]
    assert allocate(Decimal("0.02"), [Decimal("1")] * 4) == [Decimal("0"), Decimal("0"), Decimal("0"), Decimal("0.02")]
    assert allocate(Decimal("10"), [Decimal("0"), Decimal("0")]) == [Decimal("5"), Decimal("5")]
    assert sum(allocate(Decimal("6000"), [Decimal("1.5"), Decimal("2.5"), Decimal("2")])) == Decimal("6000")


async def test_day_split_in_paras_adds_up_exactly_across_entries_and_projects(db_session, make_project):
    first_project = await make_project(db_session, code="PR-A", name="A")
    second_project = await make_project(db_session, code="PR-B", name="B")
    user = _user(db_session, "trip-paras")
    worker = await _worker(db_session, regular_day_rate=100)

    entries = [
        await _entry(db_session, user, first_project, [worker], hours=2, pay_mode="full_day"),
        await _entry(db_session, user, second_project, [worker], hours=2),
        await _entry(db_session, user, first_project, [worker], hours=2),
    ]
    listed = {entry.id: entry for entry in await list_entries(None, None, None, None, db_session, user)}
    shares = [listed[entry.id].labor_amount for entry in entries]

    # Три равные трети ста динаров: остаток пары получает запись с наибольшим id.
    assert shares == [33.33, 33.33, 33.34]
    first_costs = await _costs(db_session, user, first_project)
    second_costs = await _costs(db_session, user, second_project)
    assert (first_costs.labor_amount, second_costs.labor_amount) == (66.67, 33.33)
    assert Decimal(str(first_costs.labor_amount)) + Decimal(str(second_costs.labor_amount)) == Decimal("100")


async def test_only_day_paid_workers_leave_the_hourly_part_of_the_team(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-mixed-team")
    day_worker = await _worker(db_session, "Ana")
    hourly_worker = await _worker(db_session, "Boris", regular_day_rate=800)

    entry = await _entry(
        db_session,
        user,
        project,
        [day_worker, hourly_worker],
        hours=5,
        worker_days=[WorkDiaryWorkerDayInput(worker_id=day_worker.id, pay_mode="full_day")],
    )

    # Ставка бригады 750 + 100; почасовая часть — только доля Бориса.
    assert entry.team_hourly_rate_snapshot == 850
    assert entry.hourly_labor_amount == pytest.approx(5 * 100)
    assert entry.day_labor_amount == 6000
    assert entry.labor_amount == pytest.approx(6500)
    pay = {item.worker_name: item for item in entry.worker_pay}
    assert pay["Ana"].labor_amount == 6000
    assert pay["Boris"].labor_amount == pytest.approx(500)


# --- 5. Фиксированная ставка командировки: суточные и питание не добавляются ---


async def test_fixed_trip_rate_does_not_add_per_diem_and_food(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-fixed")
    worker = await _worker(
        db_session,
        "Stevan",
        trip_pricing_mode="fixed_plus_lodging",
        trip_work_day_rate=10500,
        trip_per_diem_rate=2500,
        trip_food_rate=3000,
    )

    entry = await _entry(
        db_session,
        user,
        project,
        [worker],
        worker_days=[
            WorkDiaryWorkerDayInput(
                worker_id=worker.id,
                pay_mode="trip_day",
                per_diem_amount=2500,
                food_amount=3000,
                lodging_amount=3000,
            )
        ],
    )

    pay = entry.worker_pay[0]
    assert (pay.trip_pricing_mode, pay.per_diem_amount, pay.food_amount, pay.lodging_amount) == (
        "fixed_plus_lodging",
        0,
        0,
        3000,
    )
    assert entry.labor_amount == 10500
    assert entry.allowance_amount == 3000
    # Начисление дня совпадает с расчётом выплаты работнику за тот же день.
    payout = _calculate_payout(
        worker,
        WorkerPayoutCreate(
            worker_id=worker.id,
            payout_type="trip_final",
            date=DAY,
            period_start=DAY,
            period_end=DAY,
            lodging_amount=Decimal("3000"),
        ),
    )
    assert payout["gross_amount"] == Decimal(str(entry.labor_amount + entry.allowance_amount))


async def test_allowance_trip_rate_takes_per_diem_and_food_from_the_worker_card(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-allowances")
    worker = await _worker(
        db_session,
        "Denis",
        trip_pricing_mode="allowances",
        trip_work_day_rate=4000,
        trip_per_diem_rate=2500,
        trip_food_rate=3000,
    )

    entry = await _entry(db_session, user, project, [worker], pay_mode="trip_day")

    assert entry.labor_amount == 4000
    assert entry.allowance_amount == 5500
    assert entry.payout_amount == 9500


# --- Ставка, снимки и сверхурочные в режимах за день ---


async def test_missing_day_rate_is_flagged_instead_of_replaced(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-missing-rate")
    worker = await _worker(db_session, "Denis", trip_work_day_rate=0)

    entry = await _entry(db_session, user, project, [worker], pay_mode="trip_day")
    pay = entry.worker_pay[0]
    # Обычная дневная ставка 6000 молча не подставляется.
    assert (pay.day_rate, pay.rate_missing) == (0, True)
    assert entry.labor_amount == 0

    updated = await update_entry(
        entry.id,
        WorkDiaryEntryUpdate(
            worker_days=[WorkDiaryWorkerDayInput(worker_id=worker.id, pay_mode="trip_day", day_rate=3000)]
        ),
        db_session,
        user,
    )
    pay = updated.worker_pay[0]
    assert (pay.day_rate, pay.day_rate_manual, pay.rate_missing) == (3000, True, False)
    assert updated.labor_amount == 3000


async def test_day_rate_snapshot_survives_worker_card_changes(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-snapshot")
    worker = await _worker(db_session)
    entry = await _entry(db_session, user, project, [worker], pay_mode="full_day")

    worker.regular_day_rate = Decimal("9000")
    await db_session.commit()

    updated = await update_entry(entry.id, WorkDiaryEntryUpdate(description="Ispravka opisa"), db_session, user)
    assert updated.labor_amount == 6000
    # Та же дата и тот же режим в новой записи не перечитывают карточку.
    other = await _entry(db_session, user, project, [worker], hours=2, pay_mode="full_day")
    assert other.worker_pay[0].day_rate == 6000
    # Новый день берёт ставку из карточки заново.
    next_day = await _entry(db_session, user, project, [worker], day=NEXT_DAY, pay_mode="full_day")
    assert next_day.labor_amount == 9000


async def test_day_modes_are_a_fixed_amount_without_overtime(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-overtime")
    worker = await _worker(db_session, trip_work_day_rate=10000)

    hourly = await _entry(db_session, user, project, [worker], day=date(2026, 8, 30), hours=10)
    full_day = await _entry(db_session, user, project, [worker], day=date(2026, 8, 31), hours=10, pay_mode="full_day")
    trip_day = await _entry(db_session, user, project, [worker], hours=10, pay_mode="trip_day")

    # По часам сверх 8 часов — с коэффициентом сверхурочных, как раньше.
    assert hourly.labor_amount == pytest.approx(8 * 750 + 2 * 750 * 1.26)
    # За день платится фиксированная сумма, как и в расчёте выплат: сверхурочных нет.
    assert full_day.labor_amount == 6000
    assert trip_day.labor_amount == 10000
    assert full_day.overtime_person_hours == trip_day.overtime_person_hours == 2
    assert full_day.billable_amount == trip_day.billable_amount == 10 * 1500


async def test_worker_days_endpoint_shows_the_saved_mode_and_other_entries(db_session, make_project):
    museum = await make_project(db_session, code="PR-A", name="Muzej")
    fryer = await make_project(db_session, code="PR-B", name="Friteza")
    user = _user(db_session, "trip-worker-days")
    worker = await _worker(db_session)
    first = await _entry(db_session, user, museum, [worker], hours=1.5, pay_mode="full_day")
    second = await _entry(db_session, user, fryer, [worker], hours=4.5)

    states = await get_worker_days(DAY, str(worker.id), second.id, db_session, user)

    assert len(states) == 1
    state = states[0]
    assert (state.exists, state.pay_mode, state.day_rate, state.other_hours) == (True, "full_day", 6000, 1.5)
    assert [(item.id, item.project_name) for item in state.other_entries] == [(first.id, "Muzej")]


async def test_worker_days_must_belong_to_the_entry_workers(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-foreign-worker")
    worker = await _worker(db_session)
    stranger = await _worker(db_session, "Stranger")

    with pytest.raises(HTTPException) as error:
        await _entry(
            db_session,
            user,
            project,
            [worker],
            worker_days=[WorkDiaryWorkerDayInput(worker_id=stranger.id, pay_mode="full_day")],
        )
    assert error.value.status_code == 400

    with pytest.raises(ValidationError):
        WorkDiaryEntryCreate(
            date=DAY,
            project_id=project.id,
            worker_ids=[worker.id],
            description="Dupli",
            duration_hours=1,
            worker_days=[
                {"worker_id": worker.id, "pay_mode": "full_day"},
                {"worker_id": worker.id, "pay_mode": "hourly"},
            ],
        )


# --- 6. Аванс и окончательный расчёт не прибавляются поверх начисления ---


async def _trip_with_payouts(db, user, project, *, final=True):
    worker = await _worker(
        db,
        "Stevan",
        trip_pricing_mode="fixed_plus_lodging",
        trip_work_day_rate=10500,
        trip_advance_day_rate=6000,
    )
    await _entry(db, user, project, [worker], day=DAY, hours=6, pay_mode="trip_day", is_trip=True)
    await _entry(db, user, project, [worker], day=NEXT_DAY, hours=5, pay_mode="trip_day", is_trip=True)
    advance = await create_worker_payout(
        WorkerPayoutCreate(
            worker_id=worker.id,
            payout_type="trip_advance",
            date=DAY,
            period_start=DAY,
            period_end=NEXT_DAY,
            project_id=project.id,
        ),
        db,
        user,
    )
    final_payout = None
    if final:
        final_payout = await create_worker_payout(
            WorkerPayoutCreate(
                worker_id=worker.id,
                payout_type="trip_final",
                date=NEXT_DAY,
                period_start=DAY,
                period_end=NEXT_DAY,
                advance_paid=Decimal("12000"),
                project_id=project.id,
            ),
            db,
            user,
        )
    return worker, advance.payout, final_payout.payout if final_payout else None


async def test_advance_and_final_payout_do_not_add_to_the_diary_accrual(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-payouts")
    _, advance, final = await _trip_with_payouts(db_session, user, project)
    assert (advance.cash_paid_amount, final.cash_paid_amount) == (Decimal("12000"), Decimal("9000"))

    costs = await _costs(db_session, user, project)
    # Пока выплаты не сопоставлены с дневником, их видно отдельно — и итог двойной.
    assert costs.labor_amount == 21000
    assert costs.expenses_amount == 21000
    assert costs.other_expenses_amount == 0
    assert costs.unmatched_payout_amount == 21000
    assert costs.total_cost_amount == 42000
    assert [payout.period_accrued_amount for payout in costs.payouts] == [21000, 21000]

    for payout in (advance, final):
        await update_payout_reconciliation(
            payout.id, WorkDiaryPayoutReconciliationUpdate(reconciled=True), db_session, user
        )
    costs = await _costs(db_session, user, project)
    assert costs.total_cost_amount == 21000
    assert costs.matched_payout_amount == 21000
    assert costs.unmatched_payout_amount == 0
    assert costs.payout_excess_amount == 0
    group = costs.reconciliations[0]
    assert (group.accrued_amount, group.accrued_in_project_amount, group.paid_amount, group.balance_amount) == (
        21000,
        21000,
        21000,
        0,
    )
    assert len(group.payouts) == 2


async def test_advance_smaller_than_the_accrual_keeps_the_accrual_as_cost(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-advance-only")
    _, advance, _ = await _trip_with_payouts(db_session, user, project, final=False)

    await update_payout_reconciliation(
        advance.id, WorkDiaryPayoutReconciliationUpdate(reconciled=True), db_session, user
    )
    costs = await _costs(db_session, user, project)

    assert costs.total_cost_amount == 21000
    assert costs.matched_payout_amount == 12000
    group = costs.reconciliations[0]
    assert (group.paid_amount, group.balance_amount) == (12000, 9000)


async def test_payout_above_the_diary_accrual_keeps_the_excess_as_cost(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-excess")
    worker = await _worker(db_session, "Stevan", trip_pricing_mode="fixed_plus_lodging", trip_work_day_rate=10500)
    # В дневнике записан только первый из двух оплаченных дней.
    await _entry(db_session, user, project, [worker], day=DAY, pay_mode="trip_day")
    payout = await create_worker_payout(
        WorkerPayoutCreate(
            worker_id=worker.id,
            payout_type="trip_final",
            date=NEXT_DAY,
            period_start=DAY,
            period_end=NEXT_DAY,
            project_id=project.id,
        ),
        db_session,
        user,
    )
    await update_payout_reconciliation(
        payout.payout.id, WorkDiaryPayoutReconciliationUpdate(reconciled=True), db_session, user
    )

    costs = await _costs(db_session, user, project)
    assert costs.labor_amount == 10500
    assert costs.payout_excess_amount == 10500
    assert costs.matched_payout_amount == 10500
    assert costs.total_cost_amount == 21000


async def test_reconciliation_needs_a_period_and_rejects_overlaps(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-reconcile-rules")
    worker, advance, final = await _trip_with_payouts(db_session, user, project)
    no_period = await create_worker_payout(
        WorkerPayoutCreate(worker_id=worker.id, payout_type="regular", date=DAY, project_id=project.id, work_days=1),
        db_session,
        user,
    )
    with pytest.raises(HTTPException) as error:
        await update_payout_reconciliation(
            no_period.payout.id, WorkDiaryPayoutReconciliationUpdate(reconciled=True), db_session, user
        )
    assert error.value.status_code == 400

    overlapping = await create_worker_payout(
        WorkerPayoutCreate(
            worker_id=worker.id,
            payout_type="regular",
            date=NEXT_DAY,
            period_start=NEXT_DAY,
            period_end=NEXT_DAY,
            project_id=project.id,
            work_days=1,
        ),
        db_session,
        user,
    )
    await update_payout_reconciliation(
        advance.id, WorkDiaryPayoutReconciliationUpdate(reconciled=True), db_session, user
    )
    # Тот же период (окончательный расчёт) — одна группа сверки, пересекающийся другой — нет.
    await update_payout_reconciliation(final.id, WorkDiaryPayoutReconciliationUpdate(reconciled=True), db_session, user)
    with pytest.raises(HTTPException) as error:
        await update_payout_reconciliation(
            overlapping.payout.id, WorkDiaryPayoutReconciliationUpdate(reconciled=True), db_session, user
        )
    assert error.value.status_code == 409


async def test_moving_a_reconciled_payout_to_another_period_drops_the_link(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-reconcile-reset")
    worker, advance, _ = await _trip_with_payouts(db_session, user, project, final=False)
    await update_payout_reconciliation(
        advance.id, WorkDiaryPayoutReconciliationUpdate(reconciled=True), db_session, user
    )

    response = await update_worker_payout_link(
        advance.id,
        WorkerPayoutLinkUpdate(worker_id=worker.id, payout_type="trip_advance", period_start=DAY, period_end=DAY),
        db_session,
        user,
    )

    assert response.payout.diary_reconciled is False
    costs = await _costs(db_session, user, project)
    assert costs.unmatched_payout_amount == 12000


# --- 7. Чек на топливо не дублируется расчётным транспортом ---


async def test_fuel_receipt_is_counted_once_and_trip_services_are_billed(db_session, make_project, make_expense):
    project = await make_project(db_session)
    user = _user(db_session, "trip-fuel")
    worker = await _worker(db_session)
    await make_expense(
        db_session,
        amount=Decimal("3200"),
        status="paid",
        description="NIS gorivo",
        project_id=project.id,
        expense_date=DAY,
    )

    entry = await _entry(
        db_session,
        user,
        project,
        [worker],
        pay_mode="full_day",
        is_trip=True,
        travel_hours=3,
        travel_km=200,
        materials=[
            WorkDiaryMaterialCreate(
                description="Putni troškovi",
                source="service",
                quantity=200,
                unit="km",
                unit_price_snapshot=40,
            ),
            WorkDiaryMaterialCreate(
                description="Vreme putovanja",
                source="service",
                quantity=3,
                unit="h",
                unit_price_snapshot=1500,
            ),
        ],
    )

    assert entry.travel_km == 200
    assert entry.total_cost_amount == 6000
    assert entry.billable_amount == 6 * 1500 + 200 * 40 + 3 * 1500
    costs = await _costs(db_session, user, project)
    assert costs.person_hours == 6
    assert costs.other_expenses_amount == 3200
    assert costs.unmatched_payout_amount == 0
    assert costs.total_cost_amount == 6000 + 3200
    assert costs.margin_amount == entry.billable_amount - 9200


# --- 8. Старые записи, частичное фактурирование и Excel не меняются ---


async def test_legacy_hourly_entry_keeps_its_amounts(db_session, make_project):
    project = await make_project(db_session)
    user = _user(db_session, "trip-legacy")
    worker_a = await _worker(db_session, "Ana", regular_day_rate=800, billing_hourly_rate=300)
    worker_b = await _worker(db_session, "Boris", regular_day_rate=1200, billing_hourly_rate=450)

    entry = await _entry(
        db_session,
        user,
        project,
        [worker_a, worker_b],
        hours=10,
        per_diem=True,
        per_diem_amount=1000,
        lodging_amount=3000,
    )

    assert entry.labor_amount == 8 * 250 + 2 * 250 * 1.26
    assert entry.hourly_labor_amount == entry.labor_amount
    assert entry.day_labor_amount == 0
    assert entry.allowance_amount == 1000 * 2 + 3000
    assert entry.billable_amount == 10 * 750
    assert {item.pay_mode for item in entry.worker_pay} == {"hourly"}
    assert await _worker_day_count(db_session) == 0


async def test_cost_change_after_invoicing_keeps_the_invoice_amount(db_session, make_project):
    client = Client(name="Naručilac")
    db_session.add(client)
    await db_session.flush()
    project = await make_project(db_session)
    project.client_id = client.id
    user = _user(db_session, "trip-invoiced")
    worker = await _worker(db_session, regular_day_rate=800, billing_hourly_rate=250)

    invoiced = await _entry(db_session, user, project, [worker], hours=4)
    invoice = await create_invoice_from_entries(
        WorkDiaryInvoiceCreate(
            issued_date=DAY,
            lines=[WorkDiaryInvoiceLineCreate(entry_id=invoiced.id, name="Prvi deo", amount=Decimal("600"))],
        ),
        db_session,
        user,
    )
    assert invoiced.labor_amount == 400

    # Новая запись того же дня переводит работника на полный день: затраты
    # фактурированной записи меняются, сумма фактуры и остаток к фактурированию — нет.
    await _entry(db_session, user, project, [worker], hours=2, pay_mode="full_day")
    after = next(
        entry for entry in await list_entries(project.id, None, None, None, db_session, user) if entry.id == invoiced.id
    )

    assert after.labor_amount == 533.33
    assert (after.billable_amount, after.invoiced_amount, after.remaining_billable_amount) == (1000, 600, 400)
    assert after.billing_status == "partially_invoiced"
    income = (await db_session.execute(select(Income).where(Income.id == invoice.income_id))).scalar_one()
    assert income.amount_rsd == Decimal("600")
    with pytest.raises(HTTPException) as error:
        await update_entry(invoiced.id, WorkDiaryEntryUpdate(description="Posle fakture"), db_session, user)
    assert error.value.status_code == 409


async def test_invoice_line_includes_trip_services(db_session, make_project):
    client = Client(name="Naručilac")
    db_session.add(client)
    await db_session.flush()
    project = await make_project(db_session)
    project.client_id = client.id
    user = _user(db_session, "trip-invoice-services")
    worker = await _worker(db_session)
    entry = await _entry(
        db_session,
        user,
        project,
        [worker],
        pay_mode="full_day",
        materials=[
            WorkDiaryMaterialCreate(
                description="Troškovi izlaska",
                source="service",
                quantity=1,
                unit="usl",
                unit_price_snapshot=2500,
            )
        ],
    )

    invoice = await create_invoice_from_entries(
        WorkDiaryInvoiceCreate(
            issued_date=DAY,
            lines=[WorkDiaryInvoiceLineCreate(entry_id=entry.id, name="Radovi; Troškovi izlaska", amount=11500)],
        ),
        db_session,
        user,
    )

    assert invoice.amount_rsd == Decimal("11500")
    after = (await list_entries(project.id, None, None, None, db_session, user))[0]
    assert after.billing_status == "invoiced"


async def test_proposal_export_shows_work_and_services_but_no_internal_cost(db_session, make_project):
    project = await make_project(db_session, code="PR-2026-0100", name="Beograd")
    user = _user(db_session, "trip-export")
    worker = await _worker(db_session, regular_day_rate=7777, billing_hourly_rate=1300)
    entry = await _entry(
        db_session,
        user,
        project,
        [worker],
        pay_mode="full_day",
        is_trip=True,
        materials=[
            WorkDiaryMaterialCreate(
                description="Troškovi izlaska",
                source="service",
                quantity=1,
                unit="usl",
                unit_price_snapshot=2500,
            )
        ],
    )

    response = await export_proposal_xlsx(WorkDiaryProposalExportRequest(entry_ids=[entry.id]), db_session, user)
    content = b"".join([chunk async for chunk in response.body_iterator])
    rows = list(load_workbook(BytesIO(content), data_only=False).active.iter_rows(values_only=True))

    work_row = next(row for row in rows if str(row[2]).startswith("Radovi:"))
    service_row = next(row for row in rows if str(row[2]).startswith("Usluga: Troškovi izlaska"))
    assert (work_row[4], work_row[5], work_row[6]) == (6, 1300, 7800)
    assert service_row[6] == 2500
    values = [value for row in rows for value in row if value is not None]
    assert 7777 not in values
    assert not any("7777" in str(value) for value in values)
    assert not any(str(value).startswith("Korekcija") for value in values)
